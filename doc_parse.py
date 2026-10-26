"""Document parsing and chunking (Document Ingestion sheet, steps 2 and 4).

Primary path: Docling (layout-aware parse, tables, optional OCR) + HybridChunker (structure-aware chunks).
Fallback path: pypdf for PDFs and python-docx for .docx, used when Docling isn't installed or fails (PARSER=auto).
The fallback keeps page numbers and headings it can detect but does NOT preserve tables or run OCR; it says so in
`warnings`, so the document is marked "Needs review" rather than silently indexed at lower quality.
"""
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

import rag_config


@dataclass
class Chunk:
    index: int
    content: str                 # text shown to the agent / reviewer
    embed_text: str              # content with its heading context, what gets embedded
    section: str | None = None
    page_start: int | None = None
    page_end: int | None = None


@dataclass
class Parsed:
    chunks: list
    pages: int | None
    parser: str
    parser_version: str
    ocr_used: bool
    warnings: list = field(default_factory=list)


class ParseError(Exception):
    pass


# ---------------------------------------------------------------- Docling
def chunks_from_docling(chunker, doc) -> list:
    """Convert Docling HybridChunker output to our Chunk objects. Duck-typed so it can be unit-tested without Docling:
    needs chunker.chunk(dl_doc=...), chunker.contextualize(chunk=...), chunk.text, chunk.meta.headings,
    chunk.meta.doc_items[*].prov[*].page_no."""
    out = []
    for i, ch in enumerate(chunker.chunk(dl_doc=doc)):
        text = (ch.text or "").strip()
        if not text:
            continue
        pages = sorted({p.page_no for it in (getattr(ch.meta, "doc_items", None) or []) for p in (getattr(it, "prov", None) or [])})
        heads = getattr(ch.meta, "headings", None) or []
        out.append(Chunk(index=len(out), content=text, embed_text=chunker.contextualize(chunk=ch),
                         section=" > ".join(heads) if heads else None,
                         page_start=pages[0] if pages else None, page_end=pages[-1] if pages else None))
    return out


def _parse_docling(path: Path) -> Parsed:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling_core.transforms.chunker.hybrid_chunker import HybridChunker
    from docling_core.transforms.chunker.tokenizer.huggingface import HuggingFaceTokenizer
    from transformers import AutoTokenizer
    import importlib.metadata as md

    ocr = rag_config.ocr_enabled()
    opts = PdfPipelineOptions()
    opts.do_ocr = ocr
    opts.do_table_structure = True
    conv = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)})
    doc = conv.convert(str(path)).document
    # Token counting proxy. The default is an open tokenizer, not your embedding model's own, so chunk sizes are
    # approximate; CHUNK_MAX_TOKENS leaves headroom below the embedding model's context window.
    tok = HuggingFaceTokenizer(tokenizer=AutoTokenizer.from_pretrained(
        os.getenv("CHUNK_TOKENIZER", "sentence-transformers/all-MiniLM-L6-v2")), max_tokens=rag_config.chunk_max_tokens())
    chunks = chunks_from_docling(HybridChunker(tokenizer=tok), doc)
    warnings = []
    if not chunks:
        warnings.append("No text could be extracted." + ("" if ocr else " If this is a scanned PDF, set OCR=1 and re-index."))
    return Parsed(chunks, doc.num_pages(), "docling", md.version("docling"), ocr, warnings)


# ---------------------------------------------------------------- fallbacks
# Section-style numbering ("3 Scope", "3.1 Containment") or ALL-CAPS lines. "1. Wear gloves" (number + dot) is a list step, not a heading.
_HEADING = re.compile(r"^(?:\d+(?:\.\d+)*\s+[A-Z][^\n]{2,70}|[A-Z][A-Z0-9 /&,\-]{3,60})$")


def _is_heading(line: str) -> bool:
    line = line.strip()
    return (bool(line) and len(line) <= 90 and len(line.split()) <= 10
            and not line.endswith((".", ",", ";", ":")) and bool(_HEADING.match(line)))


def _split_long(text: str, max_chars: int) -> list:
    if len(text) <= max_chars:
        return [text]
    parts, cur = [], ""
    for sent in re.split(r"(?<=[.!?])\s+", text):
        while len(sent) > max_chars:                      # one huge "sentence": hard split
            parts.append(sent[:max_chars]); sent = sent[max_chars:]
        if cur and len(cur) + len(sent) + 1 > max_chars:
            parts.append(cur); cur = sent
        else:
            cur = f"{cur} {sent}".strip()
    if cur:
        parts.append(cur)
    return parts


def chunks_from_blocks(blocks) -> list:
    """blocks: iterable of (page or None, text, is_heading). Groups paragraphs under the latest heading into
    chunks of at most ~4 characters per token * CHUNK_MAX_TOKENS, never mixing sections."""
    max_chars = rag_config.chunk_max_tokens() * 4
    out, section, buf, pages = [], None, [], []

    def flush():
        nonlocal buf, pages
        if buf:
            text = "\n\n".join(buf)
            out.append(Chunk(index=len(out), content=text, embed_text=f"{section}\n{text}" if section else text, section=section,
                             page_start=min(pages) if pages else None, page_end=max(pages) if pages else None))
        buf, pages = [], []

    for page, text, heading in blocks:
        if heading:
            flush(); section = text; continue
        for piece in _split_long(text, max_chars):
            if buf and sum(len(b) for b in buf) + len(piece) > max_chars:
                flush()
            buf.append(piece)
            if page is not None:
                pages.append(page)
    flush()
    return out


def _parse_pypdf(path: Path, degraded: bool = True) -> Parsed:
    from pypdf import PdfReader
    import importlib.metadata as md
    reader = PdfReader(str(path))
    blocks = []
    for n, page in enumerate(reader.pages, start=1):
        para = []

        def end_para():
            if para:
                blocks.append((n, "\n".join(para), False))   # keep line breaks: "Label: value" lines stay separable
                para.clear()

        for line in (page.extract_text() or "").splitlines():
            line = line.strip()
            if not line:
                end_para()
            elif _is_heading(line):
                end_para()
                blocks.append((n, line, True))
            else:
                para.append(line)
        end_para()
    chunks = chunks_from_blocks(blocks)
    warnings = ["Parsed with the pypdf fallback: tables are not preserved as structure and scanned pages are not OCR'd."] if degraded else []
    if not chunks:
        warnings.append("No text could be extracted (scanned PDF?). Install Docling and set OCR=1, then re-index.")
    return Parsed(chunks, len(reader.pages), "pypdf", md.version("pypdf"), False, warnings)


def _parse_python_docx(path: Path, degraded: bool = True) -> Parsed:
    import docx
    import importlib.metadata as md
    d = docx.Document(str(path))
    blocks = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            blocks.append((None, t, (p.style.name or "").lower().startswith("heading") or p.style.name == "Title"))
    for table in d.tables:                                          # flatten rows so table text is at least searchable
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append((None, " | ".join(cells), False))
    chunks = chunks_from_blocks(blocks)
    warnings = ["Parsed with the python-docx fallback: page numbers are unavailable and tables are flattened to text rows."] if degraded else []
    if not chunks:
        warnings.append("No text could be extracted.")
    return Parsed(chunks, None, "python-docx", md.version("python-docx"), False, warnings)


def parse_document(path) -> Parsed:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in (".pdf", ".docx"):
        raise ParseError(f"Unsupported document type: {ext}")
    pref = rag_config.parser_pref()
    if pref not in ("auto", "docling", "pypdf"):
        raise ParseError(f"PARSER must be auto, docling or pypdf (got '{pref}').")
    note = None
    if pref in ("auto", "docling"):
        try:
            return _parse_docling(path)
        except ImportError as e:
            if pref == "docling":
                raise ParseError(f"PARSER=docling but Docling is not installed: {e}")
            note = "Docling is not installed; used the fallback parser."
        except Exception as e:                                       # Docling failed on this file
            if pref == "docling":
                raise ParseError(f"Docling failed: {type(e).__name__}: {e}")
            note = f"Docling failed ({type(e).__name__}: {str(e)[:150]}); used the fallback parser."
    try:
        degraded = pref != "pypdf"        # an explicit PARSER=pypdf is a choice, not a problem to flag
        parsed = _parse_pypdf(path, degraded) if ext == ".pdf" else _parse_python_docx(path, degraded)
    except ImportError as e:
        raise ParseError(f"No parser available for {ext}: {e}. Install docling (best) or pypdf / python-docx.")
    except Exception as e:
        raise ParseError(f"Could not read the file: {type(e).__name__}: {e}")
    if note:
        parsed.warnings.insert(0, note)
    return parsed
